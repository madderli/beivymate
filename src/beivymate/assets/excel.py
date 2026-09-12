"""One step per row. Customer headers map to stable fields; formulas are never generated."""
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

DEFAULT_MAPPING = {
    '用例编号':'number', '用例名称':'title', '用例描述':'description', '前置条件':'preconditions',
    '优先级':'priority', '步骤序号':'step_number', '步骤描述':'step_description', '期望结果':'expected_result',
    '特殊数据':'special_data', '备注':'notes', '是否已自动化':'automated', '编写人':'author',
    '更新人':'modified_by', '维护人':'maintainer', '发布状态':'publication_status', '用例状态':'lifecycle',
}
REQUIRED = set(DEFAULT_MAPPING.values()) - {'notes','publication_status','lifecycle'}


class ExcelExporter:
    def __init__(self, template: Path, mapping=None):
        self.template = Path(template)
        self.mapping = dict(DEFAULT_MAPPING if mapping is None else mapping)
        self.validate()

    def validate(self):
        wb = load_workbook(self.template)
        try:
            ws = wb.worksheets[0]
            headers = [cell.value for cell in ws[1]]
            if len(headers) != len(set(headers)) or not REQUIRED.issubset({self.mapping.get(h) for h in headers}):
                raise ValueError('Template has duplicate headers or missing required mapped columns')
            if any(field not in DEFAULT_MAPPING.values() for field in self.mapping.values()):
                raise ValueError('Unknown template field mapping')
            if ws.merged_cells.ranges:
                raise ValueError('Merged-cell templates are not supported')
        finally:
            wb.close()

    def export(self, cases, output: Path):
        wb = load_workbook(self.template)
        try:
            ws = wb.worksheets[0]
            headers = [c.value for c in ws[1]]
            styles = [copy(c._style) for c in ws[2]]
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row-1)
            for case in cases:
                for index, step in enumerate(case.steps, 1):
                    values = {**case.model_dump(), 'automated':'是' if case.automated else '否',
                              'publication_status':'已发布' if case.publication_status == 'published' else '未发布',
                              'lifecycle':{'active':'有效','retired':'废除','discarded':'撤销'}[case.lifecycle],
                              'step_number':index, 'step_description':step.description,
                              'expected_result':step.expected_result, 'special_data':step.special_data}
                    ws.append([values.get(self.mapping.get(header), '') for header in headers])
                    for column, cell in enumerate(ws[ws.max_row]):
                        if column < len(styles):
                            cell._style = copy(styles[column])
                        if isinstance(cell.value, str):
                            cell.data_type = 's'  # Prevent Excel formula injection, including customer text.
            with Path(output).open('xb') as stream:
                wb.save(stream)
        finally:
            wb.close()
