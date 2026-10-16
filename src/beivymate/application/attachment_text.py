"""Bounded, explicit text extraction for Run input snapshots; no OCR or macros."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

LIMIT = 200000


def extract(name: str, raw: bytes) -> tuple[str, str]:
    suffix = Path(name).suffix.lower()
    try:
        if suffix in {'.txt', '.md'}:
            text, method = raw.decode('utf-8-sig'), 'UTF-8 文本'
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted or len(reader.pages) > 200:
                raise ValueError('不支持加密 PDF 或超过 200 页的附件')
            pages = []
            for index, page in enumerate(reader.pages):
                value = page.extract_text() or ''
                if not value.strip():
                    raise ValueError(f'PDF 第 {index+1} 页没有可提取文字，请提供文本版本；本阶段不执行 OCR')
                pages.append(f'第 {index+1} 页\n{value}')
                if sum(map(len, pages)) > LIMIT:
                    raise ValueError('附件文字超过单次输入限制')
            text, method = '\n\n'.join(pages), 'PDF 文字层；不包含图片识别'
        elif suffix in {'.docx', '.xlsx'}:
            with ZipFile(BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('文档解压大小超过限制')
                if suffix == '.docx':
                    root = ET.fromstring(archive.read('word/document.xml'))
                    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    text = '\n'.join(''.join(p.itertext()) for p in root.findall('.//w:p', ns))
                    method = 'Word 正文与表格文字；不包含图片、页眉页脚及批注'
                else:
                    from openpyxl import load_workbook
                    book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
                    try:
                        lines = []
                        size = 0
                        for sheet in book:
                            lines.append('工作表：' + sheet.title)
                            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                                if number > 10000:
                                    raise ValueError('工作表超过 10000 行，请拆分附件')
                                line = '\t'.join('' if v is None else str(v) for v in row)
                                size += len(line)
                                if size > LIMIT:
                                    raise ValueError('附件文字超过单次输入限制')
                                if line.strip(): lines.append(f'{number}: {line}')
                        text, method = '\n'.join(lines), 'Excel 单元格文字与公式原文；不执行公式或宏，不识别图片'
                    finally:
                        book.close()
        else:
            raise ValueError('暂不支持图片识别，请提供文字版本')
        if not text.strip():
            raise ValueError('没有可提取的文字')
        if len(text) > LIMIT:
            raise ValueError('附件文字超过单次输入限制，请拆分')
        return text, method
    except Exception as exc:
        raise ValueError(f'附件 {name} 解析失败：{exc}') from exc
