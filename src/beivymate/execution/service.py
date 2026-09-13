"""One SQLite store per task; Excel files are replaceable projections of durable history."""
import json
import sqlite3
import subprocess
from contextlib import contextmanager
from pathlib import Path
from uuid import uuid4
from openpyxl import Workbook
from beivymate.model.artifact.test_execution import (
    ExecutionPlan, ExecutionItem, ExecutionArtifact, ExecutionAttempt, StepObservation, Defect, now)


class ExecutionService:
    def __init__(self, directory: Path, task_id: str):
        self.directory = Path(directory)
        self.directory.mkdir(parents=True, exist_ok=True)
        self.task_id = task_id
        self.path = self.directory/'execution.db'
        with self.db() as db:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS metadata (task TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS rounds (number INTEGER PRIMARY KEY, body TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS defects (number INTEGER PRIMARY KEY, body TEXT NOT NULL);
            ''')
            row = db.execute('SELECT task FROM metadata').fetchone()
            if row and row[0] != task_id:
                raise ValueError('Execution directory belongs to another task')
            if not row:
                db.execute('INSERT INTO metadata VALUES (?)',(task_id,))

    @contextmanager
    def db(self):
        connection=sqlite3.connect(self.path, timeout=30)
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def from_design(self, design, store, *, environment, versions, executor, modes=None, runner_ids=None):
        items=[]
        for proposal, snapshot in zip(design.proposals.cases, design.case_revisions):
            if proposal.action in {'retire','discard'}:
                continue
            case=store.latest(snapshot.id)
            if case.revision != snapshot.revision:
                raise ValueError('Design case changed; choose an explicit revision before planning')
            if case.review_status != 'accepted' or case.lifecycle != 'active':
                raise ValueError('Plan requires accepted active cases')
            case=case.model_copy(deep=True)
            case.blocking_questions=list(dict.fromkeys(case.blocking_questions+proposal.blocking_questions))
            case.requirement_refs=design.case_requirement_refs.get(case.id,case.requirement_refs)
            items.append(ExecutionItem(case=case, environment=environment, product_version=versions[case.product_id],
                executor=executor, runner_id=(runner_ids or {}).get(case.id), mode=(modes or {}).get(case.id,'automated' if case.automated else 'manual')))
        return ExecutionPlan(task_id=self.task_id,name='默认测试执行计划',source='design:'+design.id,items=items)

    def start_round(self, plan: ExecutionPlan):
        if plan.task_id != self.task_id:
            raise ValueError('Task mismatch')
        if len({item.id for item in plan.items}) != len(plan.items):
            raise ValueError('Duplicate execution item')
        keys=[(i.case.id,i.case.revision,i.environment,i.product_version) for i in plan.items]
        if len(keys)!=len(set(keys)):
            raise ValueError('Duplicate case/environment scope would overwrite the round matrix')
        for item in plan.items:
            if item.case.review_status != 'accepted' or item.case.lifecycle != 'active' or not item.case.function_id:
                raise ValueError('Only accepted active classified cases may execute')
            if item.product_version not in item.case.applicable_versions:
                raise ValueError('Case version mismatch')
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            number=db.execute('SELECT coalesce(max(number),0)+1 FROM rounds').fetchone()[0]
            artifact=ExecutionArtifact(plan=plan,round_number=number)
            self._save(db,artifact)
        self._json(plan.model_dump(mode='json'),f'plan-{number}.json')
        (self.directory/f'plan-{number}.md').write_text('# 测试执行计划\n\n```json\n'+plan.model_dump_json(indent=2)+'\n```\n',encoding='utf-8')
        self.export()
        return artifact

    def load(self, number):
        with self.db() as db:
            return self._load(db,number)

    @staticmethod
    def _load(db, number):
        row=db.execute('SELECT body FROM rounds WHERE number=?',(number,)).fetchone()
        if not row:
            raise ValueError('Round not found')
        return ExecutionArtifact.model_validate_json(row[0])

    @staticmethod
    def _save(db, artifact):
        db.execute('INSERT OR REPLACE INTO rounds VALUES (?,?)',(artifact.round_number,artifact.model_dump_json()))

    @staticmethod
    def _item(artifact, item_id):
        return next(i for i in artifact.plan.items if i.id==item_id)

    def begin(self, number, item_id, *, executor):
        if not executor.strip():
            raise ValueError('Executor required')
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            artifact=self._load(db,number)
            if artifact.completed:
                raise ValueError('Completed round is immutable; start a new round')
            item=self._item(artifact,item_id)
            previous=[a for a in artifact.attempts if a.item_id==item_id]
            if previous and previous[-1].status in {'running','paused','uncertain'}:
                raise ValueError('Reconcile or resume the existing attempt first')
            attempt=ExecutionAttempt(item_id=item_id,executor=executor,mode=item.mode)
            if item.case.blocking_questions:
                attempt.status='blocked'; attempt.reason='; '.join(item.case.blocking_questions); attempt.ended_at=now()
            artifact.attempts.append(attempt)
            self._save(db,artifact)
        self.export()
        return attempt

    def record_step(self, number, attempt_id, observation: StepObservation, *, defect_number=None):
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            artifact=self._load(db,number)
            attempt=next(a for a in artifact.attempts if a.id==attempt_id)
            item=self._item(artifact,attempt.item_id)
            if artifact.completed or attempt.mode!='manual' or attempt.status!='running':
                raise ValueError('Attempt is not accepting manual steps')
            if defect_number is not None:
                if observation.result != 'failed':
                    raise ValueError('Use record_defect_verification for retest results')
                row=db.execute('SELECT body FROM defects WHERE number=?',(defect_number,)).fetchone()
                if not row or Defect.model_validate_json(row[0]).case_id != item.case.id:
                    raise ValueError('Validate existing defect before recording: case mismatch or missing defect')
            if observation.step != len(attempt.observations)+1 or observation.step>len(item.case.steps):
                raise ValueError('Steps must be recorded in order')
            attempt.observations.append(observation)
            if observation.result!='pass' or observation.step==len(item.case.steps):
                attempt.status=observation.result; attempt.ended_at=now()
            if attempt.status=='failed':
                self._defect(db,artifact,attempt,item,defect_number)
            self._save(db,artifact)
        self.export()

    def pause_or_resume(self, number, attempt_id, *, resume=False, state_verified=False):
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            artifact=self._load(db,number)
            attempt=next(a for a in artifact.attempts if a.id==attempt_id)
            if artifact.completed or attempt.mode!='manual' or attempt.status!=('paused' if resume else 'running'):
                raise ValueError('Only an active manual attempt can pause/resume')
            if resume and not state_verified:
                raise ValueError('Verify environment and preconditions before resuming')
            attempt.status='running' if resume else 'paused'
            self._save(db,artifact)
        self.export()

    def reconcile(self, number, attempt_id, *, actor, reason):
        """Explicitly confirm an interrupted runner has stopped before allowing a whole-case retry."""
        if not actor.strip() or not reason.strip():
            raise ValueError('Reconciliation requires actor and reason')
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            artifact=self._load(db,number)
            attempt=next(a for a in artifact.attempts if a.id==attempt_id)
            if artifact.completed or attempt.status not in {'running','paused','uncertain'}:
                raise ValueError('Attempt does not require reconciliation')
            attempt.status='blocked'; attempt.reason=f'{actor}: {reason}'; attempt.ended_at=now()
            self._save(db,artifact)
        self.export()

    def run_script(self, number, item_id, *, executor, runners, authorized=False, timeout=60, defect_number=None):
        """Trusted caller supplies runner argv/cwd/script ref; never execute model-generated commands.

        Runner writes a JSON result to stdout: status pass/failed/blocked, actual,
        observed_version. Exit 0 is required; exit status alone is not a test verdict.
        """
        if not authorized:
            raise ValueError('Explicit script execution authorization required')
        artifact=self.load(number); item=self._item(artifact,item_id)
        if item.mode!='automated':
            raise ValueError('Plan selected manual execution')
        runner=runners.get(item.runner_id) if isinstance(runners,dict) else None
        if defect_number is not None:
            with self.db() as db:
                row=db.execute('SELECT body FROM defects WHERE number=?',(defect_number,)).fetchone()
                if not row or Defect.model_validate_json(row[0]).case_id != item.case.id:
                    raise ValueError('Existing defect must match the execution case')
        attempt=self.begin(number,item_id,executor=executor)
        if attempt.status=='blocked':
            return attempt
        status='blocked'; actual='Missing verified script or runner'; evidence=[]
        if runner is not None:
            log=self.directory/(attempt.id+'.log')
            try:
                if not isinstance(runner,dict) or not isinstance(runner.get('argv'),list) or not runner['argv'] or any(not isinstance(a,str) or not a for a in runner['argv']):
                    raise ValueError('Invalid runner argv')
                if not isinstance(runner.get('cwd'),str) or not Path(runner['cwd']).is_dir():
                    raise ValueError('Invalid runner directory')
                if runner.get('script_ref') not in item.case.automation_refs:
                    raise ValueError('Missing verified script reference')
                result=subprocess.run(runner['argv'],cwd=runner['cwd'],capture_output=True,text=True,timeout=timeout,check=False)
                log.write_text(result.stdout+'\n'+result.stderr,encoding='utf-8'); evidence=[str(log)]
                body=json.loads(result.stdout)
                if not isinstance(body,dict):
                    raise ValueError('Runner result must be an object')
                if result.returncode!=0 or body.get('observed_version')!=item.product_version:
                    raise ValueError('Runner error or observed version mismatch')
                status=body['status']; actual=body['actual']
                if status not in {'pass','failed','blocked'} or not isinstance(actual,str) or not actual.strip():
                    raise ValueError('Invalid runner result')
            except subprocess.TimeoutExpired:
                status='uncertain'; actual='Runner timed out; reconcile side effects before retry'
            except (OSError,ValueError,KeyError,TypeError):
                status='error'; actual='Runner failed or returned an invalid result'
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            artifact=self._load(db,number)
            saved=next(a for a in artifact.attempts if a.id==attempt.id)
            if saved.status!='running':
                raise ValueError('Attempt changed during script execution')
            saved.status=status; saved.reason=actual; saved.ended_at=now()
            saved.evidence=evidence  # Overall result: no fabricated step attribution.
            if status=='failed':
                self._defect(db,artifact,saved,item,defect_number)
            self._save(db,artifact)
        self.export()
        return saved

    def _defect(self, db, artifact, attempt, item, defect_number=None):
        if defect_number is not None:
            row=db.execute('SELECT body FROM defects WHERE number=?',(defect_number,)).fetchone()
            if not row: raise ValueError('Existing defect not found')
            defect=Defect.model_validate_json(row[0])
            if defect.case_id != item.case.id:
                raise ValueError('Explicit defect link must match the case')
            if attempt.id not in defect.attempt_ids:
                defect.attempt_ids.append(attempt.id)
                if artifact.round_number not in defect.rounds: defect.rounds.append(artifact.round_number)
                defect.evidence=list(dict.fromkeys(defect.evidence+attempt.evidence+[e for o in attempt.observations for e in o.evidence]))
                db.execute('UPDATE defects SET body=? WHERE number=?',(defect.model_dump_json(),defect_number))
            return
        # An attempt is the idempotency key. Cross-attempt merging requires explicit linking.
        for (body,) in db.execute('SELECT body FROM defects'):
            if attempt.id in Defect.model_validate_json(body).attempt_ids:
                return
        n=db.execute('SELECT coalesce(max(number),0)+1 FROM defects').fetchone()[0]
        defect=Defect(number=f'Bug-{n:05d}',title=item.case.title,
            description=f'{item.environment}; version={item.product_version}; product={item.case.product_id}; project={item.case.project_id}',
            steps=[{'expected':s.model_dump(),'observation':next((o.model_dump() for o in attempt.observations if o.step==i),None)}
                   for i,s in enumerate(item.case.steps,1)],
            actual_result=attempt.reason or '\n'.join(o.actual for o in attempt.observations),
            evidence=attempt.evidence+[e for o in attempt.observations for e in o.evidence],submitter=attempt.executor,
            case_id=item.case.id,case_revision=item.case.revision,
            requirement_refs=[r.model_dump() for r in item.case.requirement_refs],rounds=[artifact.round_number],attempt_ids=[attempt.id],
            status='registered' if artifact.plan.defect_mode=='auto' else 'draft',
            submitted_at=now() if artifact.plan.defect_mode=='auto' else None)
        db.execute('INSERT INTO defects VALUES (?,?)',(n,defect.model_dump_json()))

    def register_defect(self, number, *, actor):
        if not actor.strip(): raise ValueError('Submitter required')
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            row=db.execute('SELECT body FROM defects WHERE number=?',(number,)).fetchone()
            if not row: raise ValueError('Defect not found')
            defect=Defect.model_validate_json(row[0])
            if defect.status=='draft':
                defect.status='registered'; defect.submitter=actor; defect.submitted_at=now()
                db.execute('UPDATE defects SET body=? WHERE number=?',(defect.model_dump_json(),number))
        self.export()

    def record_defect_verification(self, defect_number, number, attempt_id, *, actor):
        if not actor.strip():
            raise ValueError('Verification actor required')
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            artifact=self._load(db,number)
            attempt=next((a for a in artifact.attempts if a.id==attempt_id),None)
            if attempt is None or attempt.status not in {'pass','failed','blocked','error'}:
                raise ValueError('Verification requires a terminal attempt')
            row=db.execute('SELECT body FROM defects WHERE number=?',(defect_number,)).fetchone()
            if not row:
                raise ValueError('Defect not found')
            defect=Defect.model_validate_json(row[0])
            if self._item(artifact,attempt.item_id).case.id != defect.case_id:
                raise ValueError('Verification case mismatch')
            if not any(v.attempt_id==attempt_id for v in defect.verifications):
                from beivymate.model.artifact.test_execution import DefectVerification
                defect.verifications.append(DefectVerification(round_number=number,attempt_id=attempt_id,
                    result=attempt.status,actor=actor))
                if number not in defect.rounds: defect.rounds.append(number)
                if attempt_id not in defect.attempt_ids: defect.attempt_ids.append(attempt_id)
                db.execute('UPDATE defects SET body=? WHERE number=?',(defect.model_dump_json(),defect_number))
        self.export()

    def finish(self, number):
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            artifact=self._load(db,number)
            if artifact.completed:
                return artifact  # Preserve the defect snapshot accepted for this round.
            for item in artifact.plan.items:
                attempts=[a for a in artifact.attempts if a.item_id==item.id]
                if not attempts or attempts[-1].status in {'running','paused','uncertain'}:
                    raise ValueError('Every item needs a terminal result before completing')
            artifact.defects=[d for (body,) in db.execute('SELECT body FROM defects ORDER BY number')
                if (d:=Defect.model_validate_json(body)) and number in d.rounds]
            artifact.deliverables={name:str(self.directory/name) for name in
                ('defects.xlsx','test-execution-results.xlsx','execution-results.json')}
            artifact.completed=True; self._save(db,artifact)
        self.export()
        return artifact

    def export(self):
        # Serialize projections with mutations so concurrent writers cannot publish stale workbooks.
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            rounds=[ExecutionArtifact.model_validate_json(r[0]) for r in db.execute('SELECT body FROM rounds ORDER BY number')]
            defects=[Defect.model_validate_json(r[0]) for r in db.execute('SELECT body FROM defects ORDER BY number')]
            wb=Workbook(); ws=wb.active; ws.title='执行结果'
            ws.append(['用例编号','修订','产品','项目','环境','版本']+[f'第{r.round_number}轮' for r in rounds])
            matrix={}
            detail=wb.create_sheet('执行明细');detail.append(['轮次','用例编号','尝试ID','执行人','方式','状态','开始','结束','实际结果','证据'])
            for index,r in enumerate(rounds):
                for item in r.plan.items:
                    key=(item.case.number,item.case.revision,item.case.product_id,item.case.project_id or '',item.environment,item.product_version)
                    values=matrix.setdefault(key,['—']*len(rounds))
                    attempts=[a for a in r.attempts if a.item_id==item.id]
                    if attempts:
                        status=attempts[-1].status
                        values[index]=status if status in {'pass','failed'} else 'blocked' if status not in {'running','paused'} else '—'
                    for a in attempts:
                        detail.append([r.round_number,item.case.number,a.id,a.executor,a.mode,a.status,a.started_at,a.ended_at,
                            a.reason+' '+json.dumps([o.model_dump() for o in a.observations],ensure_ascii=False),
                            '; '.join(a.evidence+[e for o in a.observations for e in o.evidence])])
            for key,values in matrix.items(): ws.append(list(key)+values)
            self._excel(wb,'test-execution-results.xlsx')
            wb=Workbook();ws=wb.active;ws.title='缺陷清单'
            ws.append(['缺陷编号','标题','描述','重现步骤','严重程度','优先级','补充信息','提交人','提交日期','关联用例','用例修订','关联需求','执行轮次','状态','实际结果','最近验证轮次','最近验证结果','验证历史'])
            for d in defects:
                ws.append([d.number,d.title,d.description,json.dumps(d.steps,ensure_ascii=False),d.severity,d.priority,
                    '; '.join(d.evidence),d.submitter,d.submitted_at,d.case_id,d.case_revision,
                    json.dumps(d.requirement_refs,ensure_ascii=False),','.join(map(str,d.rounds)),d.status,d.actual_result,
                    d.verifications[-1].round_number if d.verifications else None,
                    d.verifications[-1].result if d.verifications else None,
                    json.dumps([v.model_dump() for v in d.verifications],ensure_ascii=False)])
            self._excel(wb,'defects.xlsx')
            self._json({'task_id':self.task_id,'rounds':[r.model_dump(mode='json') for r in rounds],
                        'defects':[d.model_dump() for d in defects]},'execution-results.json')

    def _json(self, body, name):
        tmp=self.directory/(str(uuid4())+'.tmp')
        tmp.write_text(json.dumps(body,ensure_ascii=False,indent=2),encoding='utf-8');tmp.replace(self.directory/name)

    def _excel(self, wb, name):
        for ws in wb:
            ws.freeze_panes='A2'
            for row in ws:
                for cell in row:
                    if isinstance(cell.value,str): cell.data_type='s'
        tmp=self.directory/(str(uuid4())+'.xlsx')
        try:
            wb.save(tmp);tmp.replace(self.directory/name)
        finally:
            wb.close()
            if tmp.exists(): tmp.unlink()
