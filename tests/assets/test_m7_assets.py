from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from openpyxl import load_workbook

from beivymate.assets.cases import CaseStore
from beivymate.assets.excel import ExcelExporter, DEFAULT_MAPPING
from beivymate.model.artifact.test_design import ProductCatalog, CaseProposal

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT/'resources/template/tester/test_design/DefaultTestCaseTemplate.xlsx'


def catalog():
    return ProductCatalog(products={'p':'Func01','q':'Func02'},functions=[{'id':'pay','product_id':'p','name':'支付'}])


def proposal(**kwargs):
    data = dict(title='支付',description='验证支付',preconditions='订单存在',priority='P1',
        steps=[{'description':'=2+2','expected_result':'已支付','special_data':'测试账户'}],
        product_id='p',function_id='pay',applicable_versions=['1'],condition_refs=['condition:1'],change_reason='新增覆盖')
    data.update(kwargs)
    return CaseProposal(**data)


def test_concurrent_numbers_and_import_conflict(tmp_path):
    store = CaseStore(tmp_path/'cases.db',catalog())
    def create(_):
        return store.apply_batch([proposal()],design_id='d',actor='author',maintainer='owner')[0]
    with ThreadPoolExecutor(max_workers=4) as pool:
        cases = list(pool.map(create,range(12)))
    assert len({c.number for c in cases}) == 12
    assert sorted(c.number for c in cases)[0] == 'TC-Func01-00001'
    with pytest.raises(ValueError,match='Duplicate'):
        store.import_case(cases[0])
    with pytest.raises(ValueError):
        CaseStore(tmp_path/'cases.db',ProductCatalog(products={'new':'Func01'},functions=[]))


def test_revision_retirement_no_recycling(tmp_path):
    store = CaseStore(tmp_path/'cases.db',catalog())
    first = store.apply_batch([proposal()],design_id='d',actor='a',maintainer='m')[0]
    second = store.apply_batch([proposal(action='revise',existing_id=first.id,existing_revision=1)],design_id='e',actor='b',maintainer='x')[0]
    assert second.number == first.number and second.revision == 2 and second.author == 'a'
    assert second.modified_by == 'b' and second.maintainer == 'm' and not second.automated
    retired = store.apply_batch([proposal(action='retire',existing_id=first.id,existing_revision=2)],design_id='f',actor='b',maintainer='m')[0]
    assert retired.lifecycle == 'retired' and retired.review_status == 'draft'
    new = store.apply_batch([proposal()],design_id='g',actor='a',maintainer='m')[0]
    assert new.number == 'TC-Func01-00002'
    with pytest.raises(ValueError,match='Stale'):
        store.apply_batch([proposal(action='revise',existing_id=first.id,existing_revision=1)],design_id='h',actor='a',maintainer='m')


def test_excel_customer_mapping_steps_and_literal_text(tmp_path):
    store = CaseStore(tmp_path/'cases.db',catalog())
    case = store.apply_batch([proposal()],design_id='d',actor='a',maintainer='m')[0]
    workbook = load_workbook(TEMPLATE)
    sheet=workbook.worksheets[0]
    sheet.cell(1,1,'Case ID')
    custom=tmp_path/'custom.xlsx'; workbook.save(custom); workbook.close()
    mapping = dict(DEFAULT_MAPPING); mapping['Case ID']=mapping.pop('用例编号')
    output=tmp_path/'cases.xlsx'
    ExcelExporter(custom,mapping).export([case],output)
    workbook=load_workbook(output)
    sheet=workbook.worksheets[0]
    headers={cell.value:cell.column for cell in sheet[1]}
    assert sheet.cell(2,headers['Case ID']).value == case.number
    assert sheet.cell(2,headers['步骤描述']).value == '=2+2'
    assert sheet.cell(2,headers['步骤描述']).data_type == 's'
    assert sheet.cell(2,headers['是否已自动化']).value == '否'
    assert sheet.cell(2,headers['编写人']).value == 'a'
    workbook.close()
    with pytest.raises(FileExistsError):
        ExcelExporter(custom,mapping).export([case],output)


def test_automation_linkage_has_evidence_and_revision_history(tmp_path):
    from beivymate.runtime.memory import EvidenceReference
    store=CaseStore(tmp_path/'cases.db',catalog())
    first=store.apply_batch([proposal()],design_id='d',actor='a',maintainer='m')[0]
    evidence=tmp_path/'verification.txt'; evidence.write_text('script verified against case revision 1')
    ref=EvidenceReference.capture(evidence)
    second=store.update_automation(first.id,1,script_refs=['repo:commit:script'],evidence=[ref],actor='engineer')
    assert second.automated and second.number==first.number and not first.automated
    third=store.apply_batch([proposal(action='revise',existing_id=first.id,existing_revision=2)],design_id='e',actor='a',maintainer='m')[0]
    assert not third.automated and third.revision==3
    with pytest.raises(ValueError):
        store.update_automation(first.id,3,script_refs=['script'],evidence=[],actor='engineer')


def test_invalid_template_rejected(tmp_path):
    from openpyxl import Workbook
    wb=Workbook(); wb.active.append(['用例编号']); path=tmp_path/'invalid.xlsx'; wb.save(path)
    with pytest.raises(ValueError,match='missing'):
        ExcelExporter(path)


def test_catalog_rejects_cycle_and_cross_product_parent():
    with pytest.raises(ValueError):
        ProductCatalog(products={'p':'Func01'},functions=[{'id':'a','name':'A','product_id':'p','parent_id':'a'}])


def test_retirement_is_version_scoped_and_requires_acceptance(tmp_path):
    from types import SimpleNamespace
    store=CaseStore(tmp_path/'cases.db',catalog())
    initial=store.apply_batch([proposal(applicable_versions=['1','2'])],design_id='a',actor='a',maintainer='m')[0]
    store.accept_design(SimpleNamespace(case_revisions=[initial]))
    retired=store.apply_batch([proposal(action='retire',existing_id=initial.id,existing_revision=1,
        applicable_versions=['2'])],design_id='b',actor='a',maintainer='m')[0]
    assert store.active_for_version('p','2')
    store.accept_design(SimpleNamespace(case_revisions=[retired]))
    assert store.active_for_version('p','1')[0].number==initial.number
    assert store.active_for_version('p','2')==[]


def test_release_revision_and_unpublished_changes(tmp_path):
    from types import SimpleNamespace
    store=CaseStore(tmp_path/'cases.db',catalog())
    first=store.apply_batch([proposal()],design_id='d',actor='a',maintainer='m')[0]
    release=dict(requirement_id='R',requirement_version='r1',product_version='1',case_revisions={first.id:1},actor='release-owner')
    with pytest.raises(ValueError,match='accepted'):
        store.publish_requirement(**release)
    store.accept_design(SimpleNamespace(case_revisions=[first]))
    store.publish_requirement(**release)
    store.publish_requirement(**release)
    saved=CaseStore(store.path,catalog()).latest(first.id)
    assert saved.publication_status=='published' and len(saved.publications)==1
    assert saved.publications[0].requirement_id=='R'
    assert saved.model_dump()['publication_status']=='published'
    with pytest.raises(ValueError,match='only be retired'):
        store.apply_batch([proposal(action='discard',existing_id=first.id,existing_revision=1)],design_id='x',actor='a',maintainer='m')
    changed=store.apply_batch([proposal(action='revise',existing_id=first.id,existing_revision=1,title='修正步骤')],design_id='e',actor='b',maintainer='m')[0]
    assert changed.number==first.number and changed.publication_status=='unpublished'
    assert store.active_for_version('p','1')[0].publication_status=='published'
    with pytest.raises(ValueError,match='current'):
        store.publish_requirement(**release)
    retired=store.apply_batch([proposal(action='retire',existing_id=first.id,existing_revision=2)],design_id='f',actor='b',maintainer='m')[0]
    store.accept_design(SimpleNamespace(case_revisions=[retired]))
    assert store.active_for_version('p','1')==[]
    with store.connect() as db:
        assert db.execute('SELECT count(*) FROM publications').fetchone()[0]==1
        assert db.execute('SELECT count(*) FROM revisions').fetchone()[0]==3


def test_discard_draft_and_atomic_release_manifest(tmp_path):
    from types import SimpleNamespace
    store=CaseStore(tmp_path/'cases.db',catalog())
    first, second=store.apply_batch([proposal(),proposal()],design_id='d',actor='a',maintainer='m')
    store.accept_design(SimpleNamespace(case_revisions=[first]))
    with pytest.raises(ValueError,match='accepted'):
        store.publish_requirement(requirement_id='R',requirement_version='1',product_version='1',
            case_revisions={first.id:1,second.id:1},actor='owner')
    assert store.latest(first.id).publication_status=='unpublished'
    discarded=store.apply_batch([proposal(action='discard',existing_id=second.id,existing_revision=1)],design_id='e',actor='a',maintainer='m')[0]
    assert discarded.lifecycle=='discarded'
    store.accept_design(SimpleNamespace(case_revisions=[discarded]))
    with pytest.raises(ValueError,match='active'):
        store.publish_requirement(requirement_id='R',requirement_version='1',product_version='1',
            case_revisions={second.id:2},actor='owner')
    store.publish_requirement(requirement_id='R',requirement_version='1',product_version='1',
        case_revisions={first.id:1},actor='owner')
    assert store.latest(first.id).publication_status=='published'
    assert store.latest(second.id).publication_status=='unpublished'


def test_reuse_rejects_content_changes_preserves_original(tmp_path):
    from types import SimpleNamespace
    store=CaseStore(tmp_path/'cases.db',catalog())
    first=store.apply_batch([proposal()],design_id='d',actor='a',maintainer='m')[0]
    store.accept_design(SimpleNamespace(case_revisions=[first]))
    for changes in ({'title':'changed'}, {'applicable_versions':['2']},
                    {'steps':[{'description':'different','expected_result':'different'}]}):
        with pytest.raises(ValueError,match='use revise'):
            store.apply_batch([proposal(action='reuse',existing_id=first.id,existing_revision=1,**changes)],
                              design_id='e',actor='a',maintainer='m')
    reused=store.apply_batch([proposal(action='reuse',existing_id=first.id,existing_revision=1,
        condition_refs=['new-condition'],blocking_questions=['new question'])],design_id='e',actor='a',maintainer='m')[0]
    assert reused.revision==1 and reused.condition_refs==first.condition_refs
    assert store.latest(first.id).blocking_questions==[]
