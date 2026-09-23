import json
import sys
from pathlib import Path
import pytest
from openpyxl import load_workbook
from beivymate.execution.service import ExecutionService
from beivymate.model.artifact.test_design import CaseRevision
from beivymate.model.artifact.test_execution import ExecutionItem,ExecutionPlan,StepObservation


def plan(mode='manual'):
    case=CaseRevision(number='TC-Func01-00001',title='支付',description='验证支付',preconditions='订单存在',priority='P1',
        steps=[{'description':'支付','expected_result':'成功'},{'description':'查询','expected_result':'已支付'}],
        product_id='p',function_id='pay',applicable_versions=['1'],condition_refs=['r'],author='a',modified_by='a',maintainer='a',
        design_id='d',review_status='accepted',automation_refs=['verified-script'])
    return ExecutionPlan(task_id='T',name='计划',source='test-plan',defect_mode='auto',items=[ExecutionItem(
        case=case,environment='test',product_version='1',mode=mode,executor='tester',runner_id='local')])


def test_manual_rounds_resume_retry_and_excel(tmp_path):
    s=ExecutionService(tmp_path,'T');p=plan();r=s.start_round(p)
    a=s.begin(1,p.items[0].id,executor='human')
    s.record_step(1,a.id,StepObservation(step=1,actual='已支付',result='pass'))
    s.pause_or_resume(1,a.id)
    s=ExecutionService(tmp_path,'T')
    with pytest.raises(ValueError):s.pause_or_resume(1,a.id,resume=True)
    s.pause_or_resume(1,a.id,resume=True,state_verified=True)
    s.record_step(1,a.id,StepObservation(step=2,actual='未更新',result='failed'))
    with pytest.raises(ValueError):s.record_step(1,a.id,StepObservation(step=2,actual='补写',result='pass'))
    b=s.begin(1,p.items[0].id,executor='human')
    with pytest.raises(ValueError):s.record_step(1,b.id,StepObservation(step=2,actual='跳步',result='pass'))
    for i in (1,2):s.record_step(1,b.id,StepObservation(step=i,actual='正常',result='pass'))
    s.finish(1)
    with pytest.raises(ValueError):s.begin(1,p.items[0].id,executor='human')
    s.start_round(p);c=s.begin(2,p.items[0].id,executor='other')
    s.record_step(2,c.id,StepObservation(step=1,actual='环境不可用',result='blocked'));s.finish(2)
    wb=load_workbook(tmp_path/'test-execution-results.xlsx');row=list(wb.active.values)[1]
    assert row[-2:]==('pass','blocked');assert wb['执行明细'].max_row==4;wb.close()
    wb=load_workbook(tmp_path/'defects.xlsx');assert wb.active.max_row==2;assert wb.active['A2'].value=='Bug-00001';wb.close()
    assert s.load(1).attempts[0].status=='failed'


def test_script_and_mode_authorization(tmp_path):
    s=ExecutionService(tmp_path,'T');p=plan('automated');s.start_round(p)
    runners={'local':{'script_ref':'verified-script','cwd':str(tmp_path),'argv':[sys.executable,'-c',
        'import json; print(json.dumps({"status":"pass","actual":"assertions passed","observed_version":"1"}))']}}
    with pytest.raises(ValueError):s.run_script(1,p.items[0].id,executor='trigger',runners=runners)
    a=s.run_script(1,p.items[0].id,executor='trigger',runners=runners,authorized=True)
    assert a.status=='pass' and Path(a.evidence[0]).exists()
    s.finish(1)
    p=plan();s.start_round(p)
    with pytest.raises(ValueError,match='manual'):s.run_script(2,p.items[0].id,executor='trigger',runners=runners,authorized=True)


def test_skill_checkpoint_handoff(tmp_path):
    from beivymate.application.composition import create_tester_agent
    from beivymate.runtime.context import AgentContext
    from beivymate.model.entity.requirement import Requirement
    s=ExecutionService(tmp_path/'out','T');p=plan();s.start_round(p)
    a=s.begin(1,p.items[0].id,executor='human');s.record_step(1,a.id,StepObservation(step=1,actual='blocked',result='blocked'));s.finish(1)
    root=Path(__file__).resolve().parents[2]
    agent=create_tester_agent(str(root/'resources/configuration/workflow/m8_test_execution.md'),None,'offline',execution_service=s)
    c=AgentContext();c.set('execution_round_number',1)
    state=agent.start(Requirement(id='R',title='T',content='C'),tmp_path/'checkpoint.json',task_id='T',context=c)
    assert state.status=='waiting_review'
    state=agent.resume(tmp_path/'checkpoint.json',decision='approved',actor='reviewer',expected_subject_hash=state.subject_hash)
    assert state.status=='completed'


def test_defect_link_across_rounds_and_interruption(tmp_path):
    s=ExecutionService(tmp_path,'T');p=plan()
    for n in (1,2):
        s.start_round(p);a=s.begin(n,p.items[0].id,executor='human')
        s.record_step(n,a.id,StepObservation(step=1,actual='同一问题',result='failed'),defect_number=1 if n==2 else None)
        s.finish(n)
    report=json.loads((tmp_path/'execution-results.json').read_text())
    assert len(report['defects'])==1 and report['defects'][0]['rounds']==[1,2]
    s.start_round(p);a=s.begin(3,p.items[0].id,executor='human')
    s=ExecutionService(tmp_path,'T')
    with pytest.raises(ValueError):s.begin(3,p.items[0].id,executor='human')
    s.reconcile(3,a.id,actor='human',reason='现场已清理，可重新执行')
    b=s.begin(3,p.items[0].id,executor='human')
    assert not b.observations
    from beivymate.application.execute import load_plan
    assert load_plan(tmp_path/'plan-1.md').items[0].case.number==p.items[0].case.number


@pytest.mark.parametrize('body',['[]','null','{"status":"pass"}','{"status":"pass","actual":12,"observed_version":"1"}'])
def test_invalid_runner_results_are_terminal(tmp_path,body):
    s=ExecutionService(tmp_path,'T');p=plan('automated');s.start_round(p)
    runners={'local':{'script_ref':'verified-script','cwd':str(tmp_path),'argv':[sys.executable,'-c',f'print({body!r})']}}
    a=s.run_script(1,p.items[0].id,executor='trigger',runners=runners,authorized=True)
    assert a.status=='error' and a.ended_at
    assert s.load(1).attempts[0].status=='error'


@pytest.mark.parametrize('runner',[{}, {'script_ref':'verified-script'}, {'argv':'echo unsafe','cwd':'.'}])
def test_invalid_runner_config_persists_error(tmp_path,runner):
    s=ExecutionService(tmp_path,'T');p=plan('automated');s.start_round(p)
    a=s.run_script(1,p.items[0].id,executor='trigger',runners={'local':runner},authorized=True)
    assert a.status=='error'


def test_script_defect_reuse_and_handoff_snapshot(tmp_path):
    s=ExecutionService(tmp_path,'T');p=plan('automated')
    runners={'local':{'script_ref':'verified-script','cwd':str(tmp_path),'argv':[sys.executable,'-c',
        'import json; print(json.dumps({"status":"failed","actual":"overall assertion failed","observed_version":"1"}))']}}
    for n in (1,2):
        s.start_round(p)
        a=s.run_script(n,p.items[0].id,executor='trigger',runners=runners,authorized=True,defect_number=1 if n==2 else None)
        assert a.status=='failed' and not a.observations and a.evidence
        result=s.finish(n)
        assert len(result.defects)==1
        assert result.defects[0].actual_result=='overall assertion failed'
        assert all(step['observation'] is None for step in result.defects[0].steps)
        assert all(Path(path).exists() for path in result.deliverables.values())
    assert s.load(2).defects[0].rounds==[1,2]
    assert s.load(1).defects[0].rounds==[1]
    assert s.finish(1).defects[0].rounds==[1]
    from beivymate.runtime.context import AgentContext
    context=AgentContext();context.set('result',s.load(2))
    assert AgentContext.restore(context.snapshot()).get('result').defects[0].number=='Bug-00001'
    assert len(json.loads((tmp_path/'execution-results.json').read_text())['defects'])==1


def test_default_plan_template_is_runnable(tmp_path):
    from beivymate.application.execute import load_plan
    root=Path(__file__).resolve().parents[2]
    p=load_plan(root/'resources/skills/tester/test_execution/templates/zh-CN/DefaultExecutionPlanTemplate.md')
    assert ExecutionService(tmp_path,p.task_id).start_round(p).round_number==1
    bad=tmp_path/'bad.md';bad.write_text('# empty template')
    with pytest.raises(ValueError,match='json'):load_plan(bad)


def test_defect_prevalidation_retest_and_auto_review(tmp_path):
    from beivymate.agent.tester.skills.test_execution import TestExecutionSkill
    from beivymate.runtime.context import AgentContext
    s=ExecutionService(tmp_path,'T');p=plan();s.start_round(p)
    a=s.begin(1,p.items[0].id,executor='tester')
    observation=StepObservation(step=1,actual='failure',result='failed')
    with pytest.raises(ValueError,match='before recording'):
        s.record_step(1,a.id,observation,defect_number=999)
    assert s.load(1).attempts[0].observations==[]
    s.record_step(1,a.id,observation);first=s.finish(1)
    skill=TestExecutionSkill(s);c=AgentContext();c.set('step_id','execute');c.set('steps.execute.test_execution',first)
    assert skill.can_auto_accept(c)  # Acceptance of complete records, not release approval.
    s.start_round(p);b=s.begin(2,p.items[0].id,executor='tester')
    for i in (1,2):s.record_step(2,b.id,StepObservation(step=i,actual='fixed',result='pass'))
    s.record_defect_verification(1,2,b.id,actor='reviewer')
    s.record_defect_verification(1,2,b.id,actor='reviewer')
    second=s.finish(2)
    assert second.defects[0].rounds==[1,2]
    assert len(second.defects[0].verifications)==1
    assert second.defects[0].verifications[0].result=='pass'
    assert second.defects[0].status=='registered'
    assert s.load(1).defects[0].verifications==[]
    wb=load_workbook(tmp_path/'defects.xlsx');rows=list(wb.active.values)
    assert rows[1][rows[0].index('最近验证结果')]=='pass';wb.close()
    second.defects[0].status='draft';c.set('steps.execute.test_execution',second)
    assert not skill.can_auto_accept(c)


def test_design_runner_mapping_respects_manual_override(tmp_path):
    from types import SimpleNamespace
    s=ExecutionService(tmp_path,'T');case=plan().items[0].case
    design=SimpleNamespace(id='d',proposals=SimpleNamespace(cases=[SimpleNamespace(action='new',blocking_questions=[])]),
        case_revisions=[case],case_requirement_refs={})
    store=SimpleNamespace(latest=lambda identity:case)
    p=s.from_design(design,store,environment='test',versions={'p':'1'},executor='tester',runner_ids={case.id:'configured'})
    assert p.items[0].mode=='automated' and p.items[0].runner_id=='configured'
    p=s.from_design(design,store,environment='test',versions={'p':'1'},executor='tester',modes={case.id:'manual'})
    assert p.items[0].mode=='manual' and p.items[0].case.automated
