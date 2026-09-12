import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from beivymate.application.composition import create_tester_agent
from beivymate.assets.cases import CaseStore
from beivymate.model.artifact.test_design import ProductCatalog, DesignArtifact
from beivymate.model.artifact.requirement_understanding import UnderstandingData
from beivymate.model.artifact.test_analysis import AnalysisData
from beivymate.model.entity.requirement import Requirement
from beivymate.runtime.context import AgentContext
from beivymate.runtime.llm.gateway import LLMGateway
from beivymate.runtime.llm.models import LLMResponse

ROOT=Path(__file__).resolve().parents[2]


class Provider:
    def __init__(self, bad=False, overlap=False):
        self.calls=0
        self.bad=bad
        self.overlap=overlap

    def chat(self,request):
        self.calls+=1
        if self.calls < 3:
            cls=UnderstandingData if self.calls == 1 else AnalysisData
            data={name:{'status':'not_provided','reason':'未提供'} for name in cls.model_fields if name != 'unknowns'}
            field='objective_scope' if self.calls==1 else 'test_conditions'
            data[field]={'status':'provided','items':[{'text':'支付成功更新状态','basis':'explicit','source_refs':['requirement:R']}]}
            data['unknowns']=[]
        else:
            payload=json.loads(request.messages[1].content)
            ref=next(iter(payload['conditions']))
            data={'cases':[{'title':'支付成功','description':'验证状态','preconditions':'存在未支付订单','priority':'P1',
                'steps':[{'description':'完成支付','expected_result':'已支付','special_data':'测试订单'}],
                'product_id':'p','function_id':'pay','applicable_versions':['1'],
                'condition_refs':['invented' if self.bad else ref],'change_reason':'覆盖支付成功条件'}],
                'uncovered':[],'summary':'新增支付状态验证用例'}
        if self.calls == 3 and self.overlap:
            data['cases'][0].pop('function_id')
            data['uncovered']=[{'condition_ref':ref,'reason':'部分规则仍待确认'}]
        return LLMResponse(model='fake',content=json.dumps(data,ensure_ascii=False))


@pytest.mark.parametrize('bad,overlap',[(False,False),(True,False),(False,True),(True,True)])
def test_three_stage_pipeline(tmp_path,bad,overlap):
    catalog=ProductCatalog(products={'p':'Func01'},functions=[{'id':'pay','product_id':'p','name':'支付'}])
    store=CaseStore(tmp_path/'assets.db',catalog)
    provider=Provider(bad,overlap)
    def agent():
        return create_tester_agent(str(ROOT/'resources/configuration/workflow/m7_test_design.md'),
                                  LLMGateway(provider),'fake',design_store=store)
    context=AgentContext()
    context.set('actor','author'); context.set('maintainer','owner')
    context.set('target_versions',{'p':'1'}); context.set('design_output_directory',str(tmp_path))
    path=tmp_path/'checkpoint.json'
    state=agent().start(Requirement(id='R',title='支付',content='支付成功更新状态'),path,context=context)
    state=agent().resume(path,decision='approved',actor='reviewer',expected_subject_hash=state.subject_hash)
    if bad:
        with pytest.raises(ValueError,match='Unknown test condition'):
            agent().resume(path,decision='approved',actor='reviewer',expected_subject_hash=state.subject_hash)
        assert store.by_function('p','pay')==[]
        return
    state=agent().resume(path,decision='approved',actor='reviewer',expected_subject_hash=state.subject_hash)
    assert state.index==2 and state.status=='waiting_review' and provider.calls==3
    artifact=AgentContext.restore(state.context).get('test_design_artifact')
    assert isinstance(artifact,DesignArtifact)
    if overlap:
        assert artifact.proposals.uncovered == []
        assert '部分规则仍待确认' in artifact.case_revisions[0].blocking_questions
        assert artifact.case_revisions[0].function_id == 'pay'
        assert len(artifact.review_warnings) == 2
        assert json.loads(artifact.raw_response)['uncovered']
        assert '部分规则仍待确认' in artifact.markdown
    assert artifact.case_revisions[0].number=='TC-Func01-00001'
    assert list(artifact.coverage.values())==[['TC-Func01-00001']]
    workbook=load_workbook(tmp_path/(artifact.id+'.xlsx')); assert workbook.worksheets[0].max_row==2; workbook.close()
    state=agent().resume(path,decision='approved',actor='reviewer',expected_subject_hash=state.subject_hash)
    assert state.status=='completed' and provider.calls==3
    assert state.accepted_artifact('design','test_design').id==artifact.id
    assert store.latest(artifact.case_revisions[0].id).review_status=='accepted'


@pytest.mark.parametrize('action,target',[('reuse','1'),('revise','2')])
def test_existing_case_pipeline(tmp_path,action,target):
    from types import SimpleNamespace
    from beivymate.model.artifact.test_design import CaseProposal, CaseContent
    catalog=ProductCatalog(products={'p':'Func01'},functions=[{'id':'pay','product_id':'p','name':'支付'}])
    store=CaseStore(tmp_path/'assets.db',catalog)
    initial=CaseProposal(title='支付',description='验证状态',preconditions='订单存在',priority='P1',
        steps=[{'description':'支付','expected_result':'已支付'}],product_id='p',function_id='pay',
        applicable_versions=['1'],condition_refs=['old-condition'],change_reason='初始')
    old=store.apply_batch([initial],design_id='old',actor='author',maintainer='owner')[0]
    store.accept_design(SimpleNamespace(case_revisions=[old]))
    class ExistingProvider(Provider):
        def chat(self,request):
            if self.calls<2:
                return super().chat(request)
            self.calls+=1
            payload=json.loads(request.messages[1].content)
            assert payload['existing_cases'][0]['applicable_versions']==['1']
            content=old.model_dump(include=set(CaseContent.model_fields))
            content.update(action=action,existing_id=old.id,existing_revision=1,
                applicable_versions=[target],condition_refs=[next(iter(payload['conditions']))],
                blocking_questions=['本轮需确认'],change_reason='本次需求')
            return LLMResponse(model='fake',content=json.dumps({'cases':[content],'uncovered':[],'summary':'本次设计'},ensure_ascii=False))
    provider=ExistingProvider()
    def agent():
        return create_tester_agent(str(ROOT/'resources/configuration/workflow/m7_test_design.md'),
            LLMGateway(provider),'fake',design_store=store)
    context=AgentContext()
    for key,value in {'actor':'author','maintainer':'owner','target_versions':{'p':target},
                      'design_output_directory':str(tmp_path),'candidate_case_ids':[old.id]}.items():
        context.set(key,value)
    path=tmp_path/'checkpoint.json'
    state=agent().start(Requirement(id='R',title='支付',content='支付成功更新状态'),path,context=context)
    for _ in range(2):
        state=agent().resume(path,decision='approved',actor='reviewer',expected_subject_hash=state.subject_hash)
    artifact=AgentContext.restore(state.context).get('test_design_artifact')
    assert '本轮需确认' in artifact.markdown
    assert artifact.case_revisions[0].number==old.number
    assert artifact.case_revisions[0].revision==(1 if action=='reuse' else 2)
    assert artifact.case_revisions[0].applicable_versions==[target]
    assert list(artifact.coverage.values())==[[old.number]]
    state=agent().resume(path,decision='approved',actor='reviewer',expected_subject_hash=state.subject_hash)
    assert state.status=='completed'
    assert store.active_for_version('p','1')[0].revision==1
    if action=='revise':
        assert store.active_for_version('p','2')[0].revision==2
